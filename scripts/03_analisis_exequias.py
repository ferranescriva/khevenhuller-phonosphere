#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extracción exhaustiva del material funerario del diario de Khevenhüller,
etiquetando voz del testigo (diario) frente a aparato del editor."""

import csv, re
import extraer_referencias as m
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# Rutas: ajustar a la ubicacion local del corpus.
SRC = "../corpus/alvar_ezquerra_2015.md"
OUT_XLSX = "../data/exequias_pasajes.xlsx"
OUT_MD = "../data/exequias_dossier.md"

# ---- vocabulario funerario --------------------------------------------------
CEREMONIA = [
    "exequias", "honras", "honras fúnebres", "funeral", "funerales", "entierro",
    "entierros", "enterrar", "enterró", "enterraron", "enterrado", "enterrada",
    "sepultura", "sepultado", "sepultada", "sepelio", "sepulcro", "sepultar",
    "túmulo", "túmulos", "catafalco", "féretro", "ataúd", "andas", "mortuorio",
    "obsequias", "responso", "responsos", "oficio de difuntos", "misa de difuntos",
    "misa de réquiem", "réquiem", "officium", "cabo de año", "luto", "lutos",
    "enlutado", "enlutada", "duelo", "condolencia", "condolencias", "pésame",
    "cadáver", "difunto", "difunta", "difuntos", "cera", "hachas", "luminarias",
    "capilla ardiente", "pira",
]
MUERTE = ["muerte", "murió", "muriera", "muriese", "moría", "moriría",
          "falleció", "fallecimiento", "fallecida", "fallecido", "finó",
          "finada", "expiró"]

def build(words):
    return [(w, m.make_pattern(w)) for w in words]

RX_CER = build(CEREMONIA)
RX_MUE = build(MUERTE)

# ---- señales de voz ---------------------------------------------------------
APAR_RANGES = [(1, 8), (190, 219), (648, 654), (749, 9999)]
APPAR = re.compile(
    r"(\bags\b|\bahn\b|\bahpm\b|\bnap\b|\bbne\b|\brah\b|\bahmv\b|\bleg\.|\bfol\.|"
    r"\blibro \d|\bexp\.|\bsign\.|\bcf\.|\bvid\.|v[eé]ase|\[\d{1,3}\]|\(p\. ?\d|"
    r"\bpp?\. ?\d|ibidem|op\. ?cit|hans (se|anota|escribe|dice|apunta|recuerda))")
FIRST = re.compile(
    r"\b(asist[ií]|estuve|fui|fuimos|o[ií]mos|o[ií]|acompa[nñ][eé]|envi[eé]|"
    r"me confes[eé]|comulgu[eé]|escrib[ií]|part[ií]|llegu[eé]|llegamos|"
    r"me hall[eé]|regres[eé]|visit[eé]|tuve|mand[eé]|hice)\b")

# ---- referente (difunto) ----------------------------------------------------
DIFUNTOS = [
    (r"emperatriz", "Emperatriz María (†1603)"),
    (r"felipe (ii|segundo)|difunto rey|rey difunto|muerte del rey|rey (don )?felipe",
     "Felipe II (†1598)"),
    (r"reina (do[nñ]a )?ana|reina ana", "Reina Ana de Austria (†1580)"),
    (r"reina (do[nñ]a )?margarita", "Reina Margarita (†1611)"),
    (r"reina (do[nñ]a )?isabel|isabel de valois", "Reina Isabel de Valois"),
    (r"principe (don )?carlos|don carlos", "Príncipe don Carlos"),
    (r"emperador (maximiliano|max)|maximiliano (ii|segundo)",
     "Emperador Maximiliano II (†1576)"),
    (r"emperador (rodolfo|rudolf)|rodolfo (ii|segundo)", "Emperador Rodolfo II"),
    (r"carlos (v|quinto)|emperador carlos", "Carlos V"),
    (r"don juan de austria", "Don Juan de Austria"),
    (r"principe (don )?(fernando|diego|felipe)|principe heredero|el principe",
     "Príncipe (heredero)"),
    (r"infant[ae]", "Infante/a"),
    (r"archiduque|archiduquesa", "Archiduque/esa"),
    (r"cardenal", "Cardenal"),
    (r"papa|pont[ií]fice", "Papa/Pontífice"),
    (r"reina", "Reina (?)"),
    (r"emperador|emperatriz", "Emperador/triz (?)"),
]
DIF_RX = [(re.compile(p), lab) for p, lab in DIFUNTOS]

def page_of(off, ps, po):
    return m.page_for(off, po, ps)

def in_apar(pg):
    return any(a <= pg <= b for a, b in APAR_RANGES)

def sentence_bounds(raw, s, e, max_len=560, pad=150):
    seps = ".;:!?\n"
    left = max((raw.rfind(c, 0, s) for c in seps), default=-1)
    rights = [r for r in (raw.find(c, e) for c in seps) if r != -1]
    right = min(rights) if rights else len(raw)
    if right - left > max_len:
        left = max(left, s - pad); right = min(right, e + pad)
    return left + 1, right + 1

def difunto_of(win):
    for rx, lab in DIF_RX:
        if rx.search(win):
            return lab
    return "(general / otros)"

# ---- carga ------------------------------------------------------------------
pages = m.load_pages(SRC)
raw, page_starts = m.build_corpus(pages)
folded = m.fold(raw)
ps_offsets = [cs for cs, _ in page_starts]

# ---- recolectar coincidencias ----------------------------------------------
hits = []  # (start, end, tipo, term)
for term, rx in RX_CER:
    for mt in rx.finditer(folded):
        hits.append((mt.start(), mt.end(), "ceremonia", term))
for term, rx in RX_MUE:
    for mt in rx.finditer(folded):
        hits.append((mt.start(), mt.end(), "muerte", term))

# agrupar por frase
groups = {}
for s, e, tipo, term in hits:
    ls, rs = sentence_bounds(raw, s, e)
    key = (ls, rs)
    g = groups.setdefault(key, {"off": s, "tipos": set(), "terms": set()})
    g["off"] = min(g["off"], s)
    g["tipos"].add(tipo); g["terms"].add(term)

rows = []
for (ls, rs), g in groups.items():
    cita = re.sub(r"\s+", " ", raw[ls:rs]).strip(" .;:,–-")
    win = folded[max(0, g["off"] - 200): g["off"] + 200]
    dif = difunto_of(win)
    tipo = "ceremonia" if "ceremonia" in g["tipos"] else "muerte"
    # regla Tier2: 'muerte' pura sin referente ni ceremonia -> descartar ruido
    if tipo == "muerte" and dif == "(general / otros)":
        continue
    pg = page_of(g["off"], page_starts, ps_offsets)
    apar_local = APPAR.search(win)
    first = FIRST.search(win)
    if in_apar(pg) or apar_local:
        voz = "aparato"
        motivo = ("rango estudio/inventario" if in_apar(pg)
                  else f"señal aparato: {apar_local.group(0)}")
        if in_apar(pg) and 190 <= pg <= 219:
            motivo = "sección estudio/inventario (pp. 190-219)"
    else:
        voz = "diario"
        motivo = "cuerpo del diario"
    rows.append({
        "off": g["off"], "pagina": pg, "difunto": dif, "tipo": tipo,
        "terminos": ", ".join(sorted(g["terms"])), "voz": voz,
        "1a_persona": "sí" if first else "",
        "motivo": motivo, "cita": cita,
    })

rows.sort(key=lambda r: r["off"])
for i, r in enumerate(rows, 1):
    r["id"] = i

# ---- resumen ----------------------------------------------------------------
from collections import Counter
por_voz = Counter(r["voz"] for r in rows)
por_dif = Counter(r["difunto"] for r in rows)
por_dif_diario = Counter(r["difunto"] for r in rows if r["voz"] == "diario")
print(f"TOTAL pasajes funerarios: {len(rows)}")
print("Por voz:", dict(por_voz))
print("\nPor difunto (todos / solo diario):")
for d, n in por_dif.most_common():
    print(f"  {n:4} / {por_dif_diario.get(d,0):<4}  {d}")

# ---- Excel ------------------------------------------------------------------
HF = PatternFill("solid", fgColor="1F3864"); WB_ = Font(name="Arial", bold=True, color="FFFFFF")
wb = Workbook(); ws = wb.active; ws.title = "Exequias"
COLS = [("id","id"),("pagina","página"),("difunto","difunto/evento (probable)"),
        ("tipo","tipo"),("terminos","términos"),("voz","voz"),
        ("1a_persona","1ª persona"),("motivo","motivo clasif."),("cita","cita / contexto")]
ws.append([c for _, c in COLS])
for c in ws[1]:
    c.fill, c.font = HF, WB_
for r in rows:
    ws.append([r[k] for k, _ in COLS])
W = {"id":6,"página":8,"difunto/evento (probable)":26,"tipo":11,"términos":22,
     "voz":10,"1ª persona":10,"motivo clasif.":26,"cita / contexto":95}
from openpyxl.utils import get_column_letter
for j,(_,name) in enumerate(COLS,1):
    ws.column_dimensions[get_column_letter(j)].width = W[name]
for r in range(2, ws.max_row+1):
    for j in range(1, len(COLS)+1):
        ws.cell(r, j).font = Font(name="Arial", size=10)
cita_c = len(COLS)
for r in range(2, ws.max_row+1):
    ws.cell(r, cita_c).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

res = wb.create_sheet("Resumen")
res.append(["Difunto/evento", "pasajes (total)", "pasajes (diario)"])
for c in res[1]: c.fill, c.font = HF, WB_
for d, n in por_dif.most_common():
    res.append([d, n, por_dif_diario.get(d, 0)])
res.append([])
res.append(["Voz", "n"])
for v, n in por_voz.most_common():
    res.append([v, n])
res.column_dimensions["A"].width = 30; res.column_dimensions["B"].width = 16
res.column_dimensions["C"].width = 16
wb.save(OUT_XLSX)
print("\nExcel:", OUT_XLSX)

# ---- dossier markdown de los difuntos clave (solo voz del diario) ----------
CLAVE = ["Emperatriz María (†1603)", "Felipe II (†1598)",
         "Reina Ana de Austria (†1580)", "Emperador Maximiliano II (†1576)",
         "Príncipe don Carlos", "Don Juan de Austria"]
with open(OUT_MD, "w", encoding="utf-8") as f:
    f.write("# Dossier de exequias — diario de Khevenhüller\n\n")
    f.write("Pasajes de voz del diario (excluido el aparato del editor), "
            "agrupados por difunto. Página impresa entre corchetes.\n")
    for dif in CLAVE:
        sel = [r for r in rows if r["difunto"] == dif and r["voz"] == "diario"]
        f.write(f"\n## {dif}  ({len(sel)} pasajes)\n\n")
        for r in sel:
            marca = " · 1ª pers." if r["1a_persona"] else ""
            f.write(f"- **[p. {r['pagina']}]** ({r['tipo']}{marca}) "
                    f"_{r['terminos']}_ — {r['cita']}\n")
    # resto agrupado
    otros = sorted({r["difunto"] for r in rows} - set(CLAVE))
    f.write(f"\n## Otros referentes\n\n")
    for dif in otros:
        sel = [r for r in rows if r["difunto"] == dif and r["voz"] == "diario"]
        if sel:
            f.write(f"\n### {dif} ({len(sel)})\n\n")
            for r in sel[:40]:
                f.write(f"- **[p. {r['pagina']}]** _{r['terminos']}_ — {r['cita']}\n")
print("Dossier:", OUT_MD)
