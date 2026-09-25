#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Añade la capa de codificación (rejilla de Rostagno) al workbook base."""

from openpyxl import load_workbook
import csv
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Rutas: ajustar a la ubicacion local.
SRC = "../data/base_fonosfera.xlsx"
OUT = "../data/libro_codificacion.xlsx"

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
CODE_FILL = PatternFill("solid", fgColor="7030A0")   # cabeceras de codificación
WHITE_BOLD = Font(name=ARIAL, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Vocabularios controlados (cada uno < 255 car. para la validación)
VOCAB = {
    "Nivel (Rostagno)": ["pedal", "índice", "símbolo"],
    "Tipo": ["música", "campana", "artillería-salva", "voz-aclamación",
             "rumor-murmullo", "metáfora"],
    "Intención simbólica": ["política", "religiosa", "social", "artística", "—"],
    "Postura del testigo": ["registro neutro", "cualificación evaluativa",
                            "extrañamiento forastero", "re-significación"],
    "¿Ruido?": ["sí", "no"],
    "Subtipo rumor": ["acústico", "auditivo-social", "extra-sónico"],
}

wb = load_workbook(SRC)
ws = wb["Referencias"]

# localizar la columna 'notas' (última) para insertar las de codificación antes
header = [c.value for c in ws[1]]
notas_idx = header.index("notas") + 1          # 1-indexed
code_headers = list(VOCAB.keys()) + ["señal evaluativa (auto)"]
n_new = len(code_headers)
ws.insert_cols(notas_idx, amount=n_new)

# escribir cabeceras de codificación
first = notas_idx
for i, name in enumerate(code_headers):
    c = ws.cell(1, first + i, name)
    c.fill, c.font = CODE_FILL, WHITE_BOLD
    c.alignment = Alignment(vertical="center", wrap_text=True)

max_row = ws.max_row

# validaciones desplegables
for i, (name, opts) in enumerate(VOCAB.items()):
    col = first + i
    letter = get_column_letter(col)
    dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "Elige un valor de la lista."
    dv.errorTitle = "Valor no válido"
    dv.prompt = " / ".join(opts)
    dv.promptTitle = name
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max_row}")
    ws.column_dimensions[letter].width = 20

# fuente Arial + bordes suaves en toda la rejilla; 'cita' con ajuste de texto
cita_idx = header.index("cita / contexto") + 1
for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=ws.max_column):
    for c in row:
        if c.font is None or c.font.name != ARIAL:
            # conserva negrita/color de cabeceras ya fijadas
            if c.row == 1:
                if c.font and c.font.color and c.font.color.rgb == "00FFFFFF":
                    pass
                else:
                    c.font = WHITE_BOLD
                    c.fill = c.fill if c.fill and c.fill.fgColor.rgb != "00000000" else HEADER_FILL
            else:
                c.font = Font(name=ARIAL, size=10)
        c.border = BORDER
for r in range(2, max_row + 1):
    ws.cell(r, cita_idx).alignment = Alignment(wrap_text=True, vertical="top")

# anchos base
widths = {"id": 6, "fecha": 12, "entrada (texto)": 18, "página": 8,
          "categoría": 22, "término (lema)": 16, "detectado": 16,
          "cita / contexto": 70, "notas": 34}
for j, name in enumerate([c.value for c in ws[1]], start=1):
    if name in widths:
        ws.column_dimensions[get_column_letter(j)].width = widths[name]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# señal evaluativa automática (pre-flag del barrido de proximidad)
auto_col = first + len(VOCAB)
ws.column_dimensions[get_column_letter(auto_col)].width = 26
try:
    with open("../data/flags_evaluativos.csv", encoding="utf-8") as fh:
        flags = {int(x["id"]): x["flag"] for x in csv.DictReader(fh)}
    for rr in range(2, max_row + 1):
        rid = ws.cell(rr, 1).value
        if rid in flags:
            cc = ws.cell(rr, auto_col, flags[rid])
            cc.font = Font(name=ARIAL, size=9, color="7030A0")
            cc.alignment = Alignment(wrap_text=True, vertical="top")
except FileNotFoundError:
    pass

# ---- hoja de instrucciones (leyenda + ejemplo) -------------------------------
ins = wb.create_sheet("Instrucciones", 0)
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 26
ins.column_dimensions["B"].width = 95

def line(r, a, b, bold_a=True, fill_a=None):
    ca = ins.cell(r, 1, a); cb = ins.cell(r, 2, b)
    ca.font = Font(name=ARIAL, bold=bold_a, size=11)
    cb.font = Font(name=ARIAL, size=11)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    if fill_a:
        ca.fill = fill_a; ca.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=11)

r = 1
ins.cell(r, 1, "Codificación fonosférica — Dietario de Khevenhüller").font = Font(
    name=ARIAL, bold=True, size=14); r += 2
line(r, "Fuente", "Alvar Ezquerra (ed.), El embajador imperial Hans Khevenhüller "
     "(1538-1606) en España, BOE-MAEC, 2015. Marco: A. Rostagno, «Historical Urban "
     "Phonosphere» (2023)."); r += 2
line(r, "Cómo codificar", "Rellena las columnas moradas de la hoja «Referencias» "
     "usando los desplegables. La página remite a la impresa del libro; la fecha es "
     "orientativa (la más próxima en el texto corrido)."); r += 2

line(r, "COLUMNAS", "", fill_a=CODE_FILL); r += 1
line(r, "Nivel (Rostagno)", "pedal = fondo habitual no oído conscientemente · índice "
     "= indica sin intención de significar · símbolo = emitido con intención "
     "comunicativa. El nivel depende del sujeto: para Khevenhüller, forastero, lo que "
     "un local dejaría en pedal puede ascender a símbolo."); r += 1
line(r, "Tipo", "música / campana / artillería-salva / voz-aclamación / "
     "rumor-murmullo / metáfora."); r += 1
line(r, "Intención simbólica", "política / religiosa / social / artística / — (usa "
     "«—» cuando el nivel sea índice o pedal)."); r += 1
line(r, "Postura del testigo", "registro neutro · cualificación evaluativa (juzga la "
     "calidad: connoisseurship) · extrañamiento forastero · re-significación "
     "(affordance transferida: relee el sonido desde su habitus imperial)."); r += 1
line(r, "¿Ruido?", "sí / no. Marca lo que responde a la petición final de Rostagno: "
     "ruido y reverberación como parámetros de sentido."); r += 1
line(r, "Subtipo rumor", "Solo para la categoría «Voz y rumor». acústico (murmullo, "
     "son confuso) · auditivo-social (circula susurrado; atmósfera, Böhme/Griffero) · "
     "extra-sónico (noticia como contenido; fuera del núcleo). Ancla filológica "
     "(Covarrubias 1611): el sonido está en «murmullo» (ruido manso del agua, "
     "onomatopeya de murmur → «murmurar… medio entre dientes»); «rumor» ya es "
     "informativo en 1611. Ver hoja «Glosario Covarrubias»."); r += 2

line(r, "EJEMPLO", "", fill_a=CODE_FILL); r += 1
ex = ("p. 272 · «invitado en varias ocasiones a buenos conciertos de música»  →  "
      "Nivel: símbolo · Tipo: música · Intención: social · Postura: cualificación "
      "evaluativa («buenos») · ¿Ruido?: no · Subtipo rumor: —")
line(r, "Fila modelo", ex); r += 2
line(r, "Recuento", "La hoja «Resumen» trae los conteos por categoría y término del "
     "cribado automático (previo a tu codificación).")

# ---- hoja de glosario Covarrubias -------------------------------------------
GLOSARIO = [
    ("MURMULLO",
     "«El ruido manso que haze el agua corriente, a murmure, por la figura "
     "onomatopeya. Y de allí murmurar, que es dezir mal de alguno, medio entre "
     "dientes; y murmurador, y murmuración, de que vide supra verbo Morder.»",
     "Ancla ACÚSTICA del núcleo. El propio Covarrubias traza el paso de sonido "
     "físico → habla sub-audible («medio entre dientes»): la zona liminal, "
     "autorizada por la lexicografía de época. Codifica murmullo/murmurar como "
     "Tipo = rumor-murmullo, Subtipo = acústico."),
    ("MURMURACIÓN",
     "«Es una plática nacida de embidia, que procura manchar y obscurecer la vida "
     "y virtud agena; es un mortal veneno de la amistad… oficio de gente vil y "
     "baxa…» (con S. Agustín, S. Bernardo —la lengua murmuradora «es pinzel del "
     "demonio, y semejante a la víbora»— y Erasmo).",
     "Entrada moral: los «pecados de la lengua». Sirve para el marco social del "
     "rumor cortesano, no para lo sónico."),
    ("RUMOR",
     "«Lo que se dize, no en público, pero se esparce secretamente en el Pueblo. "
     "Lat. rumor.»",
     "En 1611 «rumor» es puramente informativo, sin componente acústico. Sus "
     "tokens tienden a Subtipo = auditivo-social o extra-sónico. El anclaje sónico "
     "lo aporta «murmullo», no esta voz."),
    ("SALVA",
     "Tres sentidos unidos por la raíz de salvo / seguridad: (1) «hazer la salva» "
     "= el maestresala (praegustator) prueba comida y bebida ante el señor, "
     "«porque da a entender que está salvo de toda traición y engaño»; (2) el "
     "disparo: «hazen salva los soldados a su Rey, à su General, y a su Capitán en "
     "ocasiones, disparando la arcabuzería por alto, y sin pelotas… en "
     "demostración de reconocimiento, paz, amistad» (lo mesmo fuertes, castillos y "
     "baxeles); (3) «salva… ó salvilla, la pieça de plata, ó oro, sobre que se "
     "sirve la copa del señor».",
     "DESAMBIGUACIÓN clave. Solo el sentido (2) es fonosférico: sonido-SÍMBOLO por "
     "antonomasia, con intención comunicativa política/social (Nivel = símbolo, "
     "Intención = política/social, ¿Ruido? = sí). El (3), la salvilla de plata, es "
     "cultura material (los «salva dorada de plata» del inventario): descártalo. "
     "El (1), ritual de corte, no es sónico. En ESTE corpus la salva sonora es "
     "rarísima: una sola «salva de la artillería»; casi todos los demás «salva» "
     "son el plato o nombres propios (San Salvador), a descartar. El paisaje de "
     "artillería viaja de hecho bajo «artillería», «disparo(s) de la artillería» "
     "y «arcabucería» (ya en «Sonido y ruido»). Matiz de nivel: esa única salva "
     "ceremonial = símbolo; los disparos sueltos suelen ser índice, salvo que el "
     "contexto marque ceremonia."),
]

glo = wb.create_sheet("Glosario Covarrubias", 1)
glo.sheet_view.showGridLines = False
glo.column_dimensions["A"].width = 22
glo.column_dimensions["B"].width = 100
glo.cell(1, 1, "Glosario — Covarrubias, Tesoro de la lengua castellana o española (1611)").font = Font(
    name=ARIAL, bold=True, size=14)
glo.cell(2, 1, "Transcripción normalizada (ſ, u/v, ç); se cita por lema (s.v.). "
         "Texto de 1611, dominio público.").font = Font(name=ARIAL, italic=True, size=10)
gr = 4
for head, transcr, uso in GLOSARIO:
    ch = glo.cell(gr, 1, head); ch.fill = CODE_FILL
    ch.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=11)
    ch.alignment = Alignment(vertical="top")
    ct = glo.cell(gr, 2, transcr)
    ct.font = Font(name=ARIAL, size=11); ct.alignment = Alignment(wrap_text=True, vertical="top")
    glo.row_dimensions[gr].height = max(28, (len(transcr) // 92 + 1) * 15)
    gr += 1
    cu1 = glo.cell(gr, 1, "uso en la codificación")
    cu1.font = Font(name=ARIAL, italic=True, size=9, color="7030A0")
    cu1.alignment = Alignment(vertical="top")
    cu2 = glo.cell(gr, 2, uso)
    cu2.font = Font(name=ARIAL, size=10, color="595959")
    cu2.alignment = Alignment(wrap_text=True, vertical="top")
    glo.row_dimensions[gr].height = max(24, (len(uso) // 92 + 1) * 14)
    gr += 2

# ---- hoja de auditoría de léxico --------------------------------------------
try:
    with open("../data/auditoria_lexico.csv", encoding="utf-8") as fh:
        arows = list(csv.reader(fh))
    aud = wb.create_sheet("Auditoría léxico")
    aud.sheet_view.showGridLines = False
    aud.cell(1, 1, "Auditoría del léxico fonosférico sobre el corpus (apariciones "
             "por término). Se conserva el instrumento completo por reproducibilidad; "
             "los términos con 0 documentan qué se buscó y no aparece.").font = Font(
        name=ARIAL, italic=True, size=10)
    aud.append([])
    aud.append(arows[0])
    hdr = aud.max_row
    for row in arows[1:]:
        aud.append([row[0], row[1], int(row[2])])
    for c in aud[hdr]:
        c.fill, c.font = HEADER_FILL, WHITE_BOLD
    aud.column_dimensions["A"].width = 30
    aud.column_dimensions["B"].width = 22
    aud.column_dimensions["C"].width = 12
    aud.freeze_panes = f"A{hdr + 1}"
    aud.auto_filter.ref = f"A{hdr}:C{aud.max_row}"
except FileNotFoundError:
    pass

wb.save(OUT)
print("Escrito:", OUT, "| filas:", max_row - 1, "| columnas:", ws.max_column)
